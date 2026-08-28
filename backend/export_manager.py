#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gestionnaire des exports (Word et PDF) pour Meeting & Project Manager AI.
"""

import os
from typing import List, Dict, Any
from datetime import datetime


class ExportManager:
    """Gestionnaire des exports de comptes-rendus."""
    
    def __init__(self):
        """Initialise le gestionnaire d'exports."""
        pass
        
    def export_to_word(self, titre: str, date: str, compte_rendu: str, 
                       taches: List[Dict[str, str]], output_path: str) -> bool:
        """
        Exporte un compte-rendu au format Word (.docx).
        
        Args:
            titre: Titre de la réunion.
            date: Date de la réunion.
            compte_rendu: Le compte-rendu formaté.
            taches: Liste des tâches.
            output_path: Chemin du fichier de sortie.
            
        Returns:
            True si l'export réussit, False sinon.
        """
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            
            doc = Document()
            
            # Titre centré
            title_para = doc.add_heading(titre, 0)
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Date en gris
            date_para = doc.add_paragraph(date)
            date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            date_run = date_para.runs[0]
            date_run.font.color.rgb = RGBColor(128, 128, 128)
            date_run.font.size = Pt(10)
            
            doc.add_paragraph()  # Espacement
            
            # Parser et ajouter le compte-rendu section par section
            self._add_formatted_content(doc, compte_rendu)
            
            # Ajouter les tâches
            if taches:
                doc.add_heading("📋 Tâches à réaliser", level=1)
                table = doc.add_table(rows=1, cols=3)
                table.style = 'Table Grid'
                
                # En-têtes
                header_cells = table.rows[0].cells
                header_cells[0].text = "Responsable"
                header_cells[1].text = "Tâche"
                header_cells[2].text = "Échéance"
                
                # Gras pour les en-têtes
                for cell in header_cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                
                # Données
                for tache in taches:
                    row = table.add_row().cells
                    row[0].text = tache.get("responsable", "Non assigné")
                    row[1].text = tache.get("tache", "")
                    row[2].text = tache.get("deadline", "À définir")
            
            # Sauvegarder
            doc.save(output_path)
            return True
            
        except Exception as e:
            print(f"Erreur lors de l'export Word: {e}")
            return False
            
    def _add_formatted_content(self, doc, content: str) -> None:
        """
        Ajoute du contenu formaté au document Word.
        
        Args:
            doc: Document python-docx.
            content: Contenu à formater.
        """
        from docx.shared import Pt
        
        lines = content.split("\n")
        current_section = None
        
        for line in lines:
            line_stripped = line.strip()
            
            if not line_stripped:
                doc.add_paragraph()
                continue
                
            # Détecter les sections avec émojis
            if line_stripped.startswith("📌"):
                doc.add_heading(line_stripped, level=1)
                current_section = "points_cles"
            elif line_stripped.startswith("⚠️"):
                doc.add_heading(line_stripped, level=1)
                current_section = "alertes"
            elif line_stripped.startswith("💡"):
                doc.add_heading(line_stripped, level=1)
                current_section = "idees"
            elif line_stripped.startswith("✅"):
                doc.add_heading(line_stripped, level=1)
                current_section = "actions"
            elif line_stripped.startswith("#"):
                # Titres Markdown
                level = min(line_stripped.count("#"), 3)
                doc.add_heading(line_stripped.lstrip("#").strip(), level=level)
            elif line_stripped.startswith("**") and "**:" in line_stripped:
                # Locuteur
                para = doc.add_paragraph()
                # Extraire le nom et le texte
                parts = line_stripped.split(":", 1)
                if len(parts) == 2:
                    speaker_run = para.add_run(parts[0].strip("*") + ":")
                    speaker_run.bold = True
                    para.add_run(parts[1])
            else:
                # Texte normal
                doc.add_paragraph(line_stripped)
                
    def export_to_pdf(self, titre: str, date: str, compte_rendu: str,
                      taches: List[Dict[str, str]], output_path: str) -> bool:
        """
        Exporte un compte-rendu au format PDF.
        
        Note: Les émojis sont remplacés par du texte car FPDF ne les supporte pas nativement.
        
        Args:
            titre: Titre de la réunion.
            date: Date de la réunion.
            compte_rendu: Le compte-rendu formaté.
            taches: Liste des tâches.
            output_path: Chemin du fichier de sortie.
            
        Returns:
            True si l'export réussit, False sinon.
        """
        try:
            from fpdf import FPDF
            
            pdf = FPDF()
            pdf.add_page()
            
            # Utiliser une police qui supporte l'UTF-8
            # Pour un support complet des caractères UTF-8, on utilise DejaVu ou similaire
            # Ici on utilise une approche simplifiée
            pdf.set_font("Helvetica", "", 12)
            
            # Titre centré
            pdf.set_font("Helvetica", "B", 16)
            pdf.cell(0, 10, titre, ln=True, align="C")
            
            # Date en gris
            pdf.set_font("Helvetica", "I", 10)
            pdf.set_text_color(128, 128, 128)
            pdf.cell(0, 10, date, ln=True, align="C")
            pdf.set_text_color(0, 0, 0)  # Reset couleur
            
            pdf.ln(10)
            
            # Remplacer les émojis par du texte
            content_text = self._replace_emojis(compte_rendu)
            
            # Ajouter le contenu ligne par ligne
            pdf.set_font("Helvetica", "", 11)
            for line in content_text.split("\n"):
                line = line.strip()
                if not line:
                    pdf.ln(5)
                    continue
                    
                # Vérifier si c'est un titre
                if line.startswith("[DECISION]") or line.startswith("[ALERTE]") or \
                   line.startswith("[IDEE]") or line.startswith("[ACTION]"):
                    pdf.set_font("Helvetica", "B", 12)
                    pdf.cell(0, 8, line, ln=True)
                    pdf.set_font("Helvetica", "", 11)
                else:
                    # Multi-cell pour gérer les longues lignes
                    pdf.multi_cell(0, 6, line)
            
            # Ajouter les tâches
            if taches:
                pdf.add_page()
                pdf.set_font("Helvetica", "B", 14)
                pdf.cell(0, 10, "TACHES A REALISER", ln=True)
                
                pdf.set_font("Helvetica", "", 10)
                y_start = pdf.get_y()
                
                # En-têtes
                pdf.set_font("Helvetica", "B", 10)
                pdf.cell(50, 8, "Responsable", border=1)
                pdf.cell(100, 8, "Tache", border=1)
                pdf.cell(40, 8, "Echeance", border=1)
                pdf.ln()
                
                # Données
                pdf.set_font("Helvetica", "", 10)
                for tache in taches:
                    responsable = tache.get("responsable", "Non assigné")[:20]
                    tache_desc = tache.get("tache", "")[:45]
                    deadline = tache.get("deadline", "A definir")[:15]
                    
                    pdf.cell(50, 8, responsable, border=1)
                    pdf.cell(100, 8, tache_desc, border=1)
                    pdf.cell(40, 8, deadline, border=1)
                    pdf.ln()
                    
                    # Nouvelle page si nécessaire
                    if pdf.get_y() > 250:
                        pdf.add_page()
            
            # Sauvegarder
            pdf.output(output_path)
            return True
            
        except Exception as e:
            print(f"Erreur lors de l'export PDF: {e}")
            return False
            
    def _replace_emojis(self, text: str) -> str:
        """
        Remplace les émojis par du texte descriptif.
        
        Args:
            text: Texte contenant des émojis.
            
        Returns:
            Texte avec les émojis remplacés.
        """
        replacements = {
            "📌": "[DECISION]",
            "⚠️": "[ALERTE]",
            "💡": "[IDEE]",
            "✅": "[ACTION]",
            "📋": "[TACHES]",
            "🎯": "[OBJECTIF]",
            "📅": "[DATE]",
            "👤": "[PERSONNE]",
        }
        
        result = text
        for emoji, replacement in replacements.items():
            result = result.replace(emoji, replacement)
            
        return result
        
    def generate_filename(self, titre: str, extension: str = "docx") -> str:
        """
        Génère un nom de fichier basé sur le titre et la date.
        
        Args:
            titre: Titre de la réunion.
            extension: Extension du fichier.
            
        Returns:
            Nom de fichier formaté.
        """
        date_str = datetime.now().strftime("%Y%m%d")
        # Nettoyer le titre pour éviter les caractères invalides
        safe_titre = "".join(c if c.isalnum() or c in " -_" else "_" for c in titre[:30])
        return f"CR_{safe_titre}_{date_str}.{extension}"
